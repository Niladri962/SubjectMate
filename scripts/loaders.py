"""Turn course files (PDF, PPTX, DOCX, IPYNB, MD, TXT, and zips of these) into LangChain Documents.

Every Document carries metadata used for citations:
    source  - file name, e.g. "L05. Stacks and recursion.pdf"
    subject - top-level folder or zip name, e.g. "data structures & algorithms"
    unit    - "p." for PDF pages, "slide" for PowerPoint, "part" for files without pages
    page    - 1-based page / slide / part number
"""
import io
import json
import re
import zipfile
from pathlib import Path

from langchain_core.documents import Document

SUPPORTED = {".pdf", ".pptx", ".docx", ".ipynb", ".md", ".txt"}
PART_CHARS = 3000  # files without pages are cut into "parts" of roughly this size for citations


def clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _junk(name: str) -> bool:
    parts = Path(name).parts
    return any(p == "__MACOSX" or p.startswith("._") or p == ".DS_Store" for p in parts)


def _safe_target(root: Path, member: str) -> Path | None:
    target = (root / member).resolve()
    return target if target.is_relative_to(root.resolve()) else None


def extract_zips(src_dir: Path, out_dir: Path) -> None:
    """Extract every zip in src_dir (and zips nested inside them) into out_dir/<zip name>/."""
    for zip_path in sorted(src_dir.rglob("*.zip")):
        _extract(zipfile.ZipFile(zip_path), out_dir / zip_path.stem)


def _extract(zf: zipfile.ZipFile, dest: Path) -> None:
    for info in zf.infolist():
        if info.is_dir() or _junk(info.filename):
            continue
        suffix = Path(info.filename).suffix.lower()
        if suffix == ".zip":
            nested = zipfile.ZipFile(io.BytesIO(zf.read(info)))
            _extract(nested, dest / Path(info.filename).with_suffix(""))
        elif suffix in SUPPORTED:
            target = _safe_target(dest, info.filename)
            if target and not (target.exists() and target.stat().st_size == info.file_size):
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(zf.read(info))


def _parts(text: str) -> list[str]:
    """Split page-less text into ~PART_CHARS pieces on paragraph boundaries."""
    parts, current = [], ""
    for para in text.split("\n\n"):
        if current and len(current) + len(para) > PART_CHARS:
            parts.append(current)
            current = ""
        current += para + "\n\n"
    if current.strip():
        parts.append(current)
    return parts


def _load_pdf(path: Path) -> list[tuple[str, int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [("p.", i, page.extract_text() or "") for i, page in enumerate(reader.pages, start=1)]


def _load_pptx(path: Path) -> list[tuple[str, int, str]]:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        out.append(("slide", i, "\n".join(texts)))
    return out


def _load_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(texts)


def _load_ipynb(path: Path) -> str:
    nb = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    cells = []
    for cell in nb.get("cells", []):
        source = "".join(cell.get("source", []))
        if source.strip():
            cells.append(source if cell.get("cell_type") == "markdown" else f"```python\n{source}\n```")
    return "\n\n".join(cells)


def load_file(path: Path, root: Path) -> tuple[list[Document], int]:
    """Return the file's non-empty pages/slides/parts and the total number of them."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        units = _load_pdf(path)
    elif suffix == ".pptx":
        units = _load_pptx(path)
    else:
        if suffix == ".docx":
            text = _load_docx(path)
        elif suffix == ".ipynb":
            text = _load_ipynb(path)
        else:
            text = path.read_text(encoding="utf-8", errors="ignore")
        units = [("part", i, t) for i, t in enumerate(_parts(clean(text)), start=1)]

    rel = path.relative_to(root)
    subject = rel.parts[0].removesuffix("_resources") if len(rel.parts) > 1 else ""
    docs = []
    for unit, number, text in units:
        text = clean(text)
        if len(text) >= 20:
            docs.append(Document(page_content=text, metadata={
                "source": path.name, "subject": subject, "unit": unit, "page": number}))
    return docs, len(units)
